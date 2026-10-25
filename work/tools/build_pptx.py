"""Сборка редактируемых слайдов ДРЕЙФ из готового HTML.
Текст и оформление читаются заново, графики вставляются из PNG.
При отсутствии графика остаётся подписанная рамка.
Запуск: work/.venv/bin/python -I work/tools/build_pptx.py
"""
import io
import re
import struct
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "out/slides/drift_portfolio_v4.html"
OUTPUT = ROOT / "out/slides/DREIF_portfolio.pptx"
CHARTS = ("price_of_urgency.png", "oneweb_model_vs_fact.png")
WIDTH, HEIGHT = 1280, 720
EMU_X, EMU_Y = Inches(13.333) / WIDTH, Inches(7.5) / HEIGHT


def norm(text):
    return re.sub(r"\s+", " ", text).strip()


def hex_color(value):
    value = value.lstrip("#")
    return "".join(c * 2 for c in value) if len(value) == 3 else value


@dataclass
class Node:
    tag: str
    attrs: dict = field(default_factory=dict)
    children: list = field(default_factory=list)
    parent: object = None

    @property
    def classes(self):
        return self.attrs.get("class", "").split()

    def nodes(self):
        return [child for child in self.children if isinstance(child, Node)]


class Document(HTMLParser):
    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.root = Node("document")
        self.stack = [self.root]
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        node = Node(tag, dict(attrs), parent=self.stack[-1])
        self.stack[-1].children.append(node)
        if tag not in {"br", "meta", "link", "img", "hr", "input", "wbr"}:
            self.stack.append(node)

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                break

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def walk(node):
    yield node
    for child in node.nodes():
        yield from walk(child)


def declarations(text):
    return dict((key.strip(), value.strip()) for key, value in
                re.findall(r"([\w-]+)\s*:\s*([^{}]+?)(?:;|$)", text))


class CSS:
    inherited = {"font-size", "font-weight", "color", "line-height",
                 "letter-spacing", "text-transform"}

    def __init__(self, text):
        self.variables = declarations(re.search(r":root\s*\{([^}]+)\}", text)[1])
        text = re.sub(r"var\((--[\w-]+)\)", lambda m: self.variables[m[1]], text)
        self.rules = []
        for selectors, body in re.findall(r"([^{}]+)\{([^{}]*)\}", text):
            for selector in selectors.split(","):
                selector = selector.strip()
                if not selector.startswith("@") and ":" not in selector:
                    rank = (selector.count("."), len(re.findall(r"(?:^|\s)[a-z][\w-]*", selector)))
                    self.rules.append((rank, selector, declarations(body)))
        self.rules.sort(key=lambda rule: rule[0])
        self.cache = {}

    @staticmethod
    def matches(node, selector):
        def simple(item, token):
            tag = re.match(r"^[a-z][\w-]*", token)
            return (not tag or item.tag == tag[0]) and all(
                cls in item.classes for cls in re.findall(r"\.([\w-]+)", token))

        tokens = selector.split()
        if not simple(node, tokens[-1]):
            return False
        ancestor = node.parent
        for token in reversed(tokens[:-1]):
            while ancestor and not simple(ancestor, token):
                ancestor = ancestor.parent
            if ancestor is None:
                return False
            ancestor = ancestor.parent
        return True

    def style(self, node):
        key = id(node)
        if key not in self.cache:
            style = {"font-size": "16px", "font-weight": "400", "color": "#14171c",
                     "line-height": "1.2", "letter-spacing": "0", "text-transform": "none"}
            if node.parent:
                style.update({k: v for k, v in self.style(node.parent).items()
                              if k in self.inherited})
            if node.tag in {"h1", "h2", "h3", "b", "strong", "th"}:
                style["font-weight"] = "700"
            for _, selector, values in self.rules:
                if self.matches(node, selector):
                    style.update(values)
            inline = re.sub(r"var\((--[\w-]+)\)", lambda m: self.variables[m[1]],
                            node.attrs.get("style", ""))
            style.update(declarations(inline))
            self.cache[key] = style
        return self.cache[key]


def number(value):
    return float(re.match(r"[-+\d.]+", str(value))[0])


def edges(style, key):
    values = [number(v) for v in style.get(key, "0").split()]
    if len(values) == 1:
        values *= 4
    elif len(values) == 2:
        values *= 2
    elif len(values) == 3:
        values.append(values[1])
    return tuple(number(style.get(f"{key}-{side}", value)) for side, value in
                 zip(("top", "right", "bottom", "left"), values))


def spacing(style):
    value = style.get("letter-spacing", "0")
    return number(value) * (number(style["font-size"]) if value.endswith("em") else 1)


class Metrics:
    """Ширины Arial из системного файла, без дополнительных зависимостей."""

    def __init__(self, bold=False):
        name = "Arial Bold.ttf" if bold else "Arial.ttf"
        win = "arialbd.ttf" if bold else "arial.ttf"
        paths = [Path("/System/Library/Fonts/Supplemental") / name,
                 Path("C:/Windows/Fonts") / win,
                 Path("/usr/share/fonts/truetype/msttcorefonts") / win]
        self.data = next((p.read_bytes() for p in paths if p.is_file()), None)
        self.widths = {}
        if self.data:
            data = self.data
            self.tables = {}
            for i in range(self.u16(4)):
                tag, _, offset, _ = struct.unpack_from(">4sIII", data, 12 + i * 16)
                self.tables[tag] = offset
            self.units = self.u16(self.tables[b"head"] + 18)
            self.count = self.u16(self.tables[b"hhea"] + 34)
            cmap = self.tables[b"cmap"]
            candidates = []
            for i in range(self.u16(cmap + 2)):
                platform, encoding, offset = struct.unpack_from(">HHI", data, cmap + 4 + i * 8)
                at = cmap + offset
                if platform in {0, 3} and self.u16(at) in {4, 12}:
                    candidates.append((self.u16(at), at))
            self.format, self.cmap = max(candidates)

    def u16(self, offset):
        return struct.unpack_from(">H", self.data, offset)[0]

    def glyph(self, code):
        at = self.cmap
        if self.format == 12:
            count = struct.unpack_from(">I", self.data, at + 12)[0]
            for i in range(count):
                start, end, glyph = struct.unpack_from(">III", self.data, at + 16 + 12 * i)
                if start <= code <= end:
                    return glyph + code - start
            return 0
        count = self.u16(at + 6) // 2
        for i in range(count):
            end = self.u16(at + 14 + i * 2)
            start = self.u16(at + 16 + count * 2 + i * 2)
            if start <= code <= end:
                delta = self.u16(at + 16 + count * 4 + i * 2)
                offset_at = at + 16 + count * 6 + i * 2
                offset = self.u16(offset_at)
                glyph = self.u16(offset_at + offset + 2 * (code - start)) if offset else code
                return (glyph + delta) % 65536 if glyph else 0
        return 0

    def width(self, char):
        if char not in self.widths:
            if self.data:
                glyph = self.glyph(ord(char))
                value = self.u16(self.tables[b"hmtx"] + 4 * min(glyph, self.count - 1)) / self.units
            else:
                value = 0.28 if char in " il.,:!|" else 0.85 if char in "ШЩЖМW@" else 0.56
            self.widths[char] = value
        return self.widths[char]


@dataclass
class Box:
    node: Node
    x: float
    y: float
    w: float
    h: float = 0
    children: list = field(default_factory=list)
    runs: list = field(default_factory=list)
    columns: list = field(default_factory=list)
    rows: list = field(default_factory=list)


class Builder:
    def __init__(self, html, charts=None, quiet=False):
        self.document = Document(html)
        css_text = "\n".join("".join(n.children) for n in walk(self.document.root) if n.tag == "style")
        self.css = CSS(css_text)
        self.metrics = [Metrics(), Metrics(True)]
        self.charts = charts if charts is not None else [SOURCE.parent / "charts" / n for n in CHARTS]
        self.quiet = quiet
        self.chart_index = 0
        self.boxes = []

    def runs(self, node):
        result = []

        def visit(item):
            if item.tag in {"svg", "style", "script"}:
                return
            if item.tag == "br":
                result.append(("\n", self.css.style(item)))
            for child in item.children:
                if isinstance(child, str):
                    result.append((re.sub(r"[ \t\r\n\f]+", " ", child), self.css.style(item)))
                else:
                    visit(child)

        visit(node)
        clean = []
        for text, style in result:
            if not clean or clean[-1][0].endswith((" ", "\n")):
                text = text.lstrip(" ")
            if text:
                clean.append((text, style))
        if clean:
            clean[-1] = (clean[-1][0].rstrip(" "), clean[-1][1])
        return clean

    def text_width(self, text, style):
        if style.get("text-transform") == "uppercase":
            text = text.upper()
        metrics = self.metrics[int(number(style["font-weight"]) >= 600)]
        return sum(metrics.width(c) for c in text) * number(style["font-size"]) + len(text) * spacing(style)

    def lines(self, runs, width, nowrap=False):
        if nowrap:
            return 1
        # Неразрывный пробел остаётся частью слова, переносы строк сохраняются.
        tokens = []
        for text, style in runs:
            for token in re.findall(r"\n| +|[^ \n]+", text):
                value = self.text_width(token, style)
                if tokens and token != "\n" and not token.startswith(" ") and tokens[-1][0] == "word":
                    tokens[-1] = ("word", tokens[-1][1] + value)
                else:
                    tokens.append(("break" if token == "\n" else "space" if token.startswith(" ") else "word", value))
        count, used, space = 1, 0, 0
        for kind, value in tokens:
            if kind == "break":
                count, used, space = count + 1, 0, 0
            elif kind == "space":
                space = value
            else:
                if used and used + space + value > width:
                    count, used = count + 1, 0
                used += (space if used else 0) + value
                space = 0
        return count

    def line_height(self, style):
        value = style["line-height"]
        return number(value) * (1 if value.endswith("px") else number(style["font-size"]))

    def layout(self, node, x, y, width):
        style = self.css.style(node)
        width = min(width, number(style.get("max-width", width)))
        box = Box(node, x, y, width)
        inline = all(child.tag in {"span", "b", "strong", "br", "em", "a"} for child in node.nodes())
        if node.tag in {"h1", "h2", "h3", "p", "li", "span", "th", "td"} or (inline and self.runs(node)):
            box.runs = self.runs(node)
            box.h = self.lines(box.runs, width - 2) * self.line_height(style)
        elif "chart" in node.classes:
            svg = next((n for n in node.nodes() if n.tag == "svg"), None)
            if svg is None:
                raise ValueError("В блоке графика не найден svg")
            box.w = min(width, number(svg.attrs["width"]))
            box.h = box.w * number(svg.attrs["height"]) / number(svg.attrs["width"])
        elif node.tag == "table":
            self.layout_table(box)
        else:
            top, right, bottom, left = edges(style, "padding")
            inner = width - left - right
            children = [n for n in node.nodes() if "foot" not in n.classes]
            if style.get("display") == "grid":
                gap = number(style.get("gap", "0"))
                tracks = style["grid-template-columns"]
                repeat = re.fullmatch(r"repeat\((\d+),\s*([^\)]+)\)", tracks)
                tracks = ([repeat[2]] * int(repeat[1])) if repeat else tracks.split()
                available = inner - gap * (len(tracks) - 1)
                fixed = sum(number(v) for v in tracks if v.endswith("px"))
                fractions = sum(number(v) for v in tracks if v.endswith("fr"))
                widths = [number(v) if v.endswith("px") else
                          (available - fixed) * number(v) / fractions for v in tracks]
                row_y = y + top
                for start in range(0, len(children), len(widths)):
                    row, col_x = [], x + left
                    for child, col_w in zip(children[start:start + len(widths)], widths):
                        row.append(self.layout(child, col_x, row_y, col_w))
                        col_x += col_w + gap
                    height = max(item.h for item in row)
                    for item in row:
                        if "card" in item.node.classes or "step" in item.node.classes:
                            item.h = height
                    box.children.extend(row)
                    row_y += height + gap
                box.h = row_y - y - (gap if children else 0) + bottom
            else:
                cursor = y + top
                for child in children:
                    margin_top, _, margin_bottom, _ = edges(self.css.style(child), "margin")
                    child_box = self.layout(child, x + left, cursor + margin_top, inner)
                    box.children.append(child_box)
                    cursor += margin_top + child_box.h + margin_bottom
                box.h = cursor - y + bottom
        return box

    def layout_table(self, box):
        rows = [n for n in walk(box.node) if n.tag == "tr"]
        count = len(rows[0].nodes())
        if any(len(row.nodes()) != count for row in rows):
            raise ValueError("В таблице неодинаковое число ячеек в строках")
        if "plan" in box.node.classes and count == 4:
            ratios = [0.135, 0.38, 0.175, 0.31]
        else:
            ratios = [1 / count] * count
        box.columns = [box.w * ratio for ratio in ratios]
        y = box.y
        for row in rows:
            cells, x = [], box.x
            height = 0
            for i, (cell, width) in enumerate(zip(row.nodes(), box.columns)):
                style = self.css.style(cell)
                top, right, bottom, left = edges(style, "padding")
                runs = self.runs(cell)
                lines = self.lines(runs, width - left - right - 2,
                                   nowrap="plan" in box.node.classes and i == 0)
                # У таблицы нужен запас под базовую линию последней строки.
                cell_height = top + bottom + lines * self.line_height(style) + number(style["font-size"]) * 0.4
                height = max(height, cell_height)
                cells.append((cell, runs))
                x += width
            box.rows.append((height, cells))
            y += height
        box.h = y - box.y

    def font(self, run, style):
        run.font.name = "Arial"
        run.font.size = Pt(number(style["font-size"]) * 0.75)
        run.font.bold = number(style["font-weight"]) >= 600
        run.font.color.rgb = RGBColor.from_string(hex_color(style["color"]))
        props = run._r.get_or_add_rPr()
        props.set("lang", "ru-RU")
        props.set("spc", str(round(spacing(style) * 75)))
        if style.get("text-transform") == "uppercase":
            props.set("cap", "all")
        if "background" in style and style["background"].startswith("#"):
            highlight = OxmlElement("a:highlight")
            color = OxmlElement("a:srgbClr")
            color.set("val", hex_color(style["background"]))
            highlight.append(color)
            props.insert_element_before(highlight, "a:latin", "a:ea", "a:cs", "a:sym")

    def text(self, frame, runs, style, bullet=False):
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(self.line_height(style) * 0.75)
        if bullet:
            props = paragraph._p.get_or_add_pPr()
            props.set("marL", "0")
            props.set("indent", str(round(-14 * EMU_X)))
            marker = OxmlElement("a:buChar")
            marker.set("char", "•")
            props.append(marker)
        for text, run_style in runs:
            for i, part in enumerate(text.split("\n")):
                if i:
                    paragraph.add_line_break()
                if part:
                    run = paragraph.add_run()
                    run.text = part
                    self.font(run, run_style)

    @staticmethod
    def position(box):
        return (round(box.x * EMU_X), round(box.y * EMU_Y),
                round(box.w * EMU_X), round(box.h * EMU_Y))

    def rectangle(self, slide, box, fill, border=None, radius=10, border_width=1):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, *self.position(box))
        if radius:
            shape.adjustments[0] = min(0.5, radius / min(box.w, box.h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(hex_color(fill))
        if border:
            shape.line.color.rgb = RGBColor.from_string(hex_color(border))
            shape.line.width = Pt(border_width * 0.75)
        else:
            shape.line.fill.background()
        return shape

    def draw(self, slide, box):
        node, style = box.node, self.css.style(box.node)
        if "card" in node.classes or "step" in node.classes:
            border = style.get("border", "")
            color = style.get("border-color") or (re.search(r"#[\da-fA-F]{6}", border)[0] if border else None)
            self.rectangle(slide, box, style.get("background", "#ffffff"), color,
                           number(style.get("border-radius", "10")),
                           number(border) if border else 1)
        if box.runs:
            shape = slide.shapes.add_textbox(*self.position(box))
            self.text(shape.text_frame, box.runs, style, node.tag == "li")
        elif "chart" in node.classes:
            if self.chart_index >= len(self.charts):
                raise ValueError("Для нового графика нужно добавить имя PNG в CHARTS")
            path = self.charts[self.chart_index]
            self.chart_index += 1
            if path.is_file():
                picture = slide.shapes.add_picture(str(path), round(box.x * EMU_X), round(box.y * EMU_Y))
                scale = min(box.w * EMU_X / picture.width, box.h * EMU_Y / picture.height)
                picture.width, picture.height = round(picture.width * scale), round(picture.height * scale)
                picture.left += round((box.w * EMU_X - picture.width) / 2)
                picture.top += round((box.h * EMU_Y - picture.height) / 2)
            else:
                if not self.quiet:
                    print(f"Нет PNG: {path.name}. Вставлена рамка, добавьте файл в out/slides/charts и повторите сборку.")
                shape = self.rectangle(slide, box, "#ffffff", self.css.variables["--line"], radius=0)
                label_style = self.css.style(Node("p", {"class": "note"}, parent=node))
                self.text(shape.text_frame, [(f"Нет PNG: {path.name}", label_style)], label_style)
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        elif node.tag == "table":
            self.draw_table(slide, box)
        for child in box.children:
            self.draw(slide, child)

    def draw_table(self, slide, box):
        table = slide.shapes.add_table(len(box.rows), len(box.columns), *self.position(box)).table
        for column, width in zip(table.columns, box.columns):
            column.width = round(width * EMU_X)
        for row_index, (height, cells) in enumerate(box.rows):
            table.rows[row_index].height = round(height * EMU_Y)
            for col_index, (node, runs) in enumerate(cells):
                cell = table.cell(row_index, col_index)
                style = self.css.style(node)
                top, right, bottom, left = edges(style, "padding")
                self.text(cell.text_frame, runs, style)
                cell.margin_top, cell.margin_bottom = round(top * EMU_Y), round(bottom * EMU_Y)
                cell.margin_left, cell.margin_right = round(left * EMU_X), round(right * EMU_X)
                cell.vertical_anchor = MSO_ANCHOR.TOP
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
                props = cell._tc.get_or_add_tcPr()
                for border_index, side in enumerate(("L", "R", "T", "B")):
                    line = OxmlElement(f"a:ln{side}")
                    line.set("w", str(round(0.75 * 12700)))
                    fill = OxmlElement("a:solidFill" if side == "B" else "a:noFill")
                    if side == "B":
                        color = OxmlElement("a:srgbClr")
                        color.set("val", self.css.variables["--line"].lstrip("#"))
                        fill.append(color)
                    line.append(fill)
                    props.insert(border_index, line)
        table.first_row = table.horz_banding = table.vert_banding = False

    def footer(self, slide, node):
        style = self.css.style(node)
        left = number(style["left"])
        width = WIDTH - left - number(style["right"])
        spans = node.nodes()
        if len(spans) != 2:
            raise ValueError("В подвале ожидаются источник и номер слайда")
        right_width = sum(self.text_width(t, s) for t, s in self.runs(spans[1])) + 3
        gap = number(style["gap"])
        first = self.layout(spans[0], left, 0, width - gap - right_width)
        second = self.layout(spans[1], left + width - right_width, 0, right_width)
        height = max(first.h, second.h)
        y = HEIGHT - number(style["bottom"]) - height
        first.y = second.y = y
        line = Box(node, left, y - number(style["padding-top"]), width, 0.5)
        self.rectangle(slide, line, self.css.variables["--line"], radius=0)
        self.draw(slide, first)
        self.draw(slide, second)
        return line.y

    def build(self):
        sections = [n for n in walk(self.document.root) if n.tag == "section" and "slide" in n.classes]
        if len(sections) != 5:
            raise ValueError(f"Ожидалось 5 слайдов, найдено {len(sections)}")
        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
        props = presentation.core_properties
        props.author = props.last_modified_by = "DevCore"
        props.title = "ДРЕЙФ"
        props.subject = props.comments = props.keywords = props.category = ""
        props.created = props.modified = datetime.now(timezone.utc)
        for i, section in enumerate(sections, 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
            box = self.layout(section, 0, 0, WIDTH)
            self.boxes.append(box)
            self.draw(slide, box)
            foot = next(n for n in section.nodes() if "foot" in n.classes)
            limit = self.footer(slide, foot)
            bottom = max(child.y + child.h for child in box.children)
            if bottom > limit - 4:
                raise ValueError(f"Слайд {i}: содержание до {bottom:.1f} px, подвал с {limit:.1f} px. Уменьшите объём или измените CSS.")
        return presentation


class VisibleText(HTMLParser):
    """Независимое извлечение фрагментов для проверки готового файла."""

    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.fragments = []
        self.slide = None
        self.excluded = 0
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        if tag in {"svg", "style", "script", "head"}:
            self.excluded += 1
        if tag == "section" and "slide" in dict(attrs).get("class", "").split():
            self.fragments.append([])
            self.slide = self.fragments[-1]

    def handle_endtag(self, tag):
        if tag in {"svg", "style", "script", "head"}:
            self.excluded -= 1
        if tag == "section":
            self.slide = None

    def handle_data(self, data):
        if self.slide is not None and not self.excluded and norm(data):
            self.slide.append(norm(data))


def missing_text(html, presentation):
    missing = []
    sources = VisibleText(html).fragments
    for index, fragments in enumerate(sources):
        texts = []
        if index < len(presentation.slides):
            for shape in presentation.slides[index].shapes:
                if shape.has_table:
                    texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
                elif shape.has_text_frame:
                    texts.append(shape.text)
        target, offset = norm(" ".join(texts)), 0
        for fragment in fragments:
            found = target.find(fragment, offset)
            if found < 0:
                missing.append(f"Слайд {index + 1}: {fragment}")
            else:
                offset = found + len(fragment)
    return missing


def package_bytes(presentation):
    raw, result = io.BytesIO(), io.BytesIO()
    presentation.save(raw)
    with ZipFile(raw) as source, ZipFile(result, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            if entry.filename == "docProps/thumbnail.jpeg":
                continue
            data = source.read(entry)
            if entry.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship\b[^>]*Type="[^"]*/thumbnail"[^>]*/>', b"", data)
            elif entry.filename == "docProps/app.xml":
                ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
                ET.register_namespace("", ns)
                root = ET.fromstring(data)
                for name, value in {"Application": "", "AppVersion": "", "Company": "DevCore",
                                    "PresentationFormat": "16:9", "Slides": "5"}.items():
                    root.find(f"{{{ns}}}{name}").text = value
                for name in ("HeadingPairs", "TitlesOfParts"):
                    element = root.find(f"{{{ns}}}{name}")
                    if element is not None:
                        root.remove(element)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, data)
    return result.getvalue()


def main():
    html = SOURCE.read_text(encoding="utf-8")
    builder = Builder(html)
    presentation = builder.build()
    data = package_bytes(presentation)
    restored = Presentation(io.BytesIO(data))
    missing = missing_text(html, restored)
    print(f"потеряно фрагментов: {len(missing)}")
    for fragment in missing:
        print(fragment)
    if missing:
        raise ValueError("Текст потерян, итоговый файл не перезаписан")
    OUTPUT.write_bytes(data)
    print(f"Готово: {OUTPUT.relative_to(ROOT)}, слайдов: {len(restored.slides)}")


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError) as error:
        print(f"Ошибка сборки: {error}", file=sys.stderr)
        sys.exit(1)
